"""Relatório «Contas Convênio» (formulário em blocos) — fixtures SINTÉTICAS que imitam o layout real:
colunas B/C/F/G/J/K/N/P/Q/T, título repetido, total ajustado à mão, anotações soltas, rodapé e uma
aba-resumo com outra versão do custo. Nenhum dado real (LGPD)."""
import io

import openpyxl
import pytest

from analise_planilha import config, contas_convenio as cc, leitura, servico
from analise_planilha.modelo import carregar_modelo

MODELO = config.PASTA_MODELOS / "contas_convenio.toml"


def _bloco(ws, linha, cod, paciente, convenio, proc, grupos, total, custo=None, res=None,
           cabecalho_custo=False, anotacao=None):
    """Escreve um bloco de atendimento a partir de `linha`; devolve a próxima linha livre."""
    ws[f"B{linha}"], ws[f"C{linha}"], ws[f"F{linha}"], ws[f"G{linha}"] = "Cód. Atd.:", cod, "Paciente:", paciente
    ws[f"B{linha + 2}"], ws[f"C{linha + 2}"], ws[f"K{linha + 2}"] = "Convênio:", convenio, proc
    h = linha + 5
    ws[f"B{h}"], ws[f"J{h}"], ws[f"P{h}"] = "Grupo Faturamento", "Qtd.", "Valor"
    if cabecalho_custo:
        ws[f"Q{h}"], ws[f"T{h}"] = "CUSTO", "RESULTADO"
    r = h + 3
    for nome, qtd, valor in grupos:
        ws[f"B{r}"], ws[f"J{r}"], ws[f"P{r}"] = nome, str(qtd), valor     # Qtd. vem como TEXTO
        r += 2
    ws[f"N{r}"], ws[f"P{r}"] = "Valor total:", total
    if custo is not None:
        ws[f"Q{r}"], ws[f"T{r}"] = custo, res
    if anotacao is not None:
        ws[f"V{r}"] = anotacao                                           # conta manual ao lado
    return r + 3


def relatorio(com_resumo=True, rodape=True) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Page 1"
    ws["H2"] = "Contas Convênio"
    lin = 5
    # 1) cesariana com custo apurado e TOTAL AJUSTADO À MÃO (1500 ≠ 1000+500+800)
    lin = _bloco(ws, lin, "100001", "MARIA SINTETICA UM", "BRADESCO SAÚDE S.A.", "CESARIANA (FETO UNICO OU MULTIPLO)",
                 [("DIARIAS", 1, 1000.0), ("Mat/Med", 12, 500.0), ("HONORARIOS MEDICOS", 1, 800.0)],
                 total=1500.0, custo=1000.0, res=500.0, cabecalho_custo=True, anotacao=123.45)
    ws[f"U{lin - 1}"] = 999.0                                           # célula solta entre blocos
    ws[f"H{lin}"] = "Contas Convênio"                                   # título de página repetido
    # 2) internação sem custo: total = soma dos grupos
    lin = _bloco(ws, lin + 2, "100002", "JOSE SINTETICO DOIS", "BRADESCO SAUDE - OPERADORA DE PLANOS S/A",
                 "INTERNACAO CLINICA", [("DIARIAS", 3, 3000.0), ("MAT ESPECIAIS, ORTESES E PROTESES (OPME)", 1, 2000.0)],
                 total=5000.0)
    # 3) cirurgia, outro convênio, sem custo
    lin = _bloco(ws, lin, "100003", "ANA SINTETICA TRES", "AMIL ASSISTENCIA MEDICA INTERNACIONAL S/A",
                 "POSTECTOMIA", [("PACOTES ESPECIAIS", 1, 4000.0), ("TAXAS", 2, 250.0)], total=4250.0)
    if rodape:
        ws[f"P{lin}"], ws[f"Q{lin}"], ws[f"T{lin}"] = 1500.0 + 5000.0 + 4250.0, 1000.0, 500.0
    if com_resumo:
        rs = wb.create_sheet("CESARIANA (FETO ÚNICO OU MÚLTIP")
        rs["A1"] = "1. CESARIANA (FETO ÚNICO OU MÚLTIPLO)"
        rs.append(["Paciente", "Faturamento (R$)", "Custo (R$)", "Resultado (R$)"])
        rs.append(["MARIA SINTETICA UM", 1500.0, 900.0, 600.0])          # custo diverge da Page 1
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture
def rel_2025() -> bytes:
    return relatorio()


# ---- leitura ---------------------------------------------------------------
def test_reconhece_e_le_uma_linha_por_atendimento(rel_2025):
    df, meta = leitura.ler(rel_2025, "2025.xlsx")
    assert meta.aba == "Page 1"                                   # aba do relatório, não a do resumo
    assert len(df) == 3
    assert list(df.index) == [5, 24, 39]                          # linha real do «Cód. Atd.» (rastreio)
    assert set(df[cc.C_ANO]) == {2025}
    assert df[cc.C_CONV].tolist()[2] == "AMIL ASSISTENCIA MEDICA INTERNACIONAL S/A"
    assert df.loc[5, "Mat/Med"] == 500.0


def test_bruto_e_apurado_nao_se_misturam(rel_2025):
    df, _ = leitura.ler(rel_2025, "2025.xlsx")
    ces, inter = df.loc[5], df.loc[24]
    assert ces[cc.C_BRUTO] == 2300.0 and ces[cc.C_APURADO] == 1500.0         # total foi ajustado à mão
    assert ces[cc.C_AJUSTADO] == "Sim" and ces[cc.C_TEM_CUSTO] == "Sim"
    assert ces[cc.C_MARGEM] == pytest.approx(500 / 1500)
    assert inter[cc.C_BRUTO] == 5000.0 and inter[cc.C_TOTAL_REL] == 5000.0
    assert inter[cc.C_TEM_CUSTO] == "Não" and inter[cc.C_AJUSTADO] == "Não"
    assert inter[cc.C_APURADO] != inter[cc.C_APURADO]                       # NaN: sem apuração
    assert df.loc[24, cc.C_DIARIAS] == 3.0                                   # Qtd. em texto → número


def test_diagnostico_transparente(rel_2025):
    _, meta = leitura.ler(rel_2025, "2025.xlsx")
    texto = " ".join(meta.diagnostico)
    assert "3 atendimentos" in texto and "Ano 2025" in texto
    assert "Custo apurado em 1 de 3" in texto
    assert "1 «Valor total» foram ajustados à mão" in texto
    assert "ignorado" in texto                                     # anotações manuais contadas
    assert "rodapé confere" in texto
    assert "Divergências" in texto and "custo em 1" in texto       # aba-resumo com outro custo


def test_aba_resumo_vira_colunas(rel_2025):
    df, _ = leitura.ler(rel_2025, "2025.xlsx")
    assert df.loc[5, cc.C_RES_CUSTO] == 900.0 and df.loc[5, cc.C_RES_RES] == 600.0
    assert df[cc.C_RES_CUSTO].notna().sum() == 1                   # só a cesariana correspondente


def test_sem_resumo_nem_rodape():
    df, meta = leitura.ler(relatorio(com_resumo=False, rodape=False), "contas.xlsx")
    assert cc.C_RES_CUSTO not in df.columns
    assert df[cc.C_ANO].isna().all()                               # sem ano no nome do arquivo
    assert not any("rodapé" in d for d in meta.diagnostico)


def test_planilha_comum_nao_e_confundida(xlsx_com_titulo):
    _, meta = leitura.ler(xlsx_com_titulo, "comum.xlsx")
    assert not any("Contas Convênio" in d for d in meta.diagnostico)


# ---- consolidação e modelo -------------------------------------------------
@pytest.fixture
def modelo():
    return carregar_modelo(MODELO)


def test_modelo_aplica_nomes_e_regras(rel_2025, modelo):
    ds = servico.carregar(rel_2025, "2025.xlsx", modelo=modelo)
    assert {"Materiais e medicamentos", "OPME", "Tipo de atendimento"} <= set(ds.df.columns)
    assert ds.colunas_faltando == []                               # grupos ausentes são opcionais
    assert set(ds.df["Convênio"]) == {"BRADESCO SAÚDE", "AMIL"}     # razões sociais unificadas
    assert dict(zip(ds.df["Procedimento"], ds.df["Tipo de atendimento"]))["INTERNACAO CLINICA"] == "Internação clínica"


def test_consolida_anos_e_indicadores_financeiros(modelo):
    ds = servico.carregar_varios([(relatorio(), "2025.xlsx"), (relatorio(com_resumo=False), "2026.xlsx")],
                                 modelo=modelo)
    assert len(ds.df) == 6 and set(ds.df["Ano"].astype(str)) == {"2025", "2026"}
    assert set(ds.df[config.COL_ARQUIVO]) == {"2025.xlsx", "2026.xlsx"}
    an = servico.analisar(ds)
    kpis = {k.rotulo: k for k in an.kpis}
    assert kpis["Margem (contas apuradas)"].valor == "33,3%"       # 2×500 / 2×1500, ponderada
    assert kpis["Contas com custo apurado"].valor == "33,3%"
    assert "em 2 de 6" in kpis["Custo (soma)"].detalhe             # cobertura parcial explícita
    ano = next(p for p in an.paineis if p.aba == "Ano").tabela
    assert {"Faturamento bruto (soma)", "Faturamento bruto (média)", "Custo (soma)", "Resultado (soma)"} <= set(ano.columns)
    comp = next(p for p in an.paineis if p.aba == "Composição").tabela
    assert comp["% de Faturamento bruto"].sum() == pytest.approx(1.0)
    cruz = [p.titulo for p in an.paineis if p.aba == "Cruzamentos"]
    assert "Tipo de atendimento por ano" in cruz


def test_sem_avisos_de_vazio_esperado(rel_2025, modelo):
    an = servico.analisar(servico.carregar(rel_2025, "2025.xlsx", modelo=modelo))
    vazios = [a.texto for a in an.alertas if "valores vazios" in a.texto]
    assert not any("OPME" in t or "Custo" in t for t in vazios)    # vazio por regra de negócio


def test_qualidade_aponta_arquivo_na_consolidacao(modelo):
    a = relatorio()
    b = relatorio(com_resumo=False)
    ds = servico.carregar_varios([(a, "2025.xlsx"), (b, "2026.xlsx")], modelo=modelo)
    # mesmo «Cód. Atendimento» nos dois arquivos → duplicidade, com a origem de cada ocorrência
    dup = ds.problemas[ds.problemas["Problema"].str.contains("uplic", na=False)]
    assert not dup.empty and dup["Aba"].str.contains("2026.xlsx").any()
