"""Leitura estrutural de planilhas "do mundo real" — fixtures SINTÉTICAS (nunca dados reais — LGPD)."""
import datetime as dt
import io

import openpyxl
import pandas as pd
import pytest

from analise_planilha import config, servico
from analise_planilha.modelo import ColunaModelo, Modelo
from analise_planilha.modelos import DATA, MOEDA


def _xlsx(abas: dict, mesclas: dict | None = None, ocultas=()) -> bytes:
    """abas: {nome: [linhas]} ; mesclas: {nome: ["A1:C1", ...]}."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for nome, linhas in abas.items():
        ws = wb.create_sheet(nome)
        for linha in linhas:
            ws.append(list(linha))
        for faixa in (mesclas or {}).get(nome, []):
            ws.merge_cells(faixa)
        if nome in ocultas:
            ws.sheet_state = "hidden"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _tipos(ds):
    return {c.nome: c.tipo for c in ds.esquema.colunas}


def _diag(ds) -> str:
    return " | ".join(ds.meta.diagnostico)


# ---------------------------------------------------------------------------
def test_relatorio_gerencial_completo():
    """Título, metadados, cabeçalho em 2 níveis mesclado, seções, subtotais, total geral,
    erro de fórmula e nota de rodapé — tudo na mesma aba."""
    linhas = [
        ["RELATÓRIO DE ATENDIMENTOS — 1º TRIMESTRE"],
        ["Unidade:", "Matriz"],
        [],
        ["Identificação", None, "Financeiro", None, None],
        ["Data", "Cidade", "Qtd", "Valor", "Glosa"],
        ["REGIÃO SERRANA"],
        [dt.datetime(2026, 1, 5), "Petrópolis", 3, 120.5, 0],
        [dt.datetime(2026, 1, 6), "Teresópolis", 2, 80.0, 10],
        ["Subtotal", None, 5, 200.5, 10],
        ["REGIÃO METROPOLITANA"],
        [dt.datetime(2026, 1, 7), "Niterói", 4, 300.0, "#DIV/0!"],
        [dt.datetime(2026, 1, 8), "Rio", 1, 50.0, 0],
        ["Subtotal", None, 5, 350.0, 0],
        ["TOTAL GERAL", None, 10, 550.5, 10],
        [],
        ["Fonte: sistema interno. Valores em R$."],
    ]
    dados = _xlsx({"Atendimentos": linhas}, {"Atendimentos": ["A1:E1", "A4:B4", "C4:E4", "A6:E6", "A10:E10"]})
    ds = servico.carregar(dados, "gerencial.xlsx")
    df = ds.df

    assert ds.meta.linhas_cabecalho == [3, 4]
    assert list(servico.visiveis(df).columns)[:6] == [
        "Seção", "Identificação - Data", "Identificação - Cidade",
        "Financeiro - Qtd", "Financeiro - Valor", "Financeiro - Glosa"]
    assert len(df) == 4                                   # só os 4 atendimentos
    assert df["Financeiro - Valor"].sum() == pytest.approx(550.5)  # nada somado em dobro
    assert list(df["Seção"]) == ["REGIÃO SERRANA"] * 2 + ["REGIÃO METROPOLITANA"] * 2
    assert list(df[config.COL_LINHA]) == [7, 8, 11, 12]  # nº real da linha no Excel
    assert ds.meta.descartes["total/subtotal"] == 3
    assert ds.meta.descartes["nota/rodapé"] == 1
    assert "erro de fórmula" in _diag(ds)
    assert _tipos(ds)["Identificação - Data"] == DATA


def test_cabecalho_reimpresso_a_cada_pagina_csv():
    cab = "Data;Cidade;Valor"
    blocos = []
    for p in range(3):
        blocos.append(cab)
        blocos += [f"0{p + 1}/02/2026;Rio;{10 * (i + 1)},00" for i in range(5)]
        blocos.append("")
    ds = servico.carregar("\n".join(blocos).encode(), "paginado.csv")
    assert len(ds.df) == 15
    assert ds.meta.descartes["cabeçalho repetido"] == 2
    assert _tipos(ds)["Valor"] == MOEDA


def test_csv_com_preambulo_nao_perde_linhas():
    """Antes, linhas com mais campos que a 1ª linha eram descartadas em silêncio."""
    texto = "Exportado pelo sistema X\nPeríodo: 01/2026\n\nNome;Cidade;Valor\n"
    texto += "\n".join(f"Pessoa {i};Rio;{i},50" for i in range(1, 41))
    ds = servico.carregar(texto.encode("cp1252"), "export.csv")
    assert ds.meta.separador == ";"
    assert len(ds.df) == 40
    assert ds.meta.linha_cabecalho == 3


def test_escolhe_a_aba_com_os_dados():
    instrucoes = [["Como preencher"], ["1. Use uma linha por atendimento"], ["2. Não apague colunas"]]
    base = [["Data", "Cidade", "Valor"]] + [[dt.datetime(2026, 3, d), "Rio", d * 10.0] for d in range(1, 21)]
    oculta = [["Código", "Descrição"]] + [[i, f"item {i}"] for i in range(200)]
    dados = _xlsx({"Instruções": instrucoes, "Base": base, "Tabelas": oculta}, ocultas=["Tabelas"])
    ds = servico.carregar(dados, "multi.xlsx")
    assert ds.meta.aba == "Base"
    assert len(ds.df) == 20
    assert "escolhida automaticamente" in _diag(ds)


def test_modelo_guia_o_cabecalho_e_associa_por_semelhanca():
    linhas = [
        ["Controle de beneficiários"],
        ["Responsável", "Fulano", "Setor", "Recepção"],
        ["Nome do Beneficiário", "Cidade", "Vlr. Total Pago", "Situação"],
    ] + [[f"Pessoa {i}", "Rio", f"{i},00", "Ativo"] for i in range(1, 11)]
    modelo = Modelo("Teste", [
        ColunaModelo("nome", "Beneficiário", "texto_livre", ["nome do beneficiario"]),
        ColunaModelo("valor", "Valor total", "moeda", ["valor total pago"]),
    ])
    ds = servico.carregar(_xlsx({"Plan1": linhas}), "benef.xlsx", modelo=modelo)
    assert ds.meta.linha_cabecalho == 2
    assert "Valor total" in ds.df.columns            # "Vlr. Total Pago" ≈ "valor total pago"
    assert _tipos(ds)["Valor total"] == MOEDA
    assert "por semelhança" in _diag(ds)
    assert ds.colunas_faltando == []


def test_segunda_tabela_abaixo_e_ignorada():
    linhas = [["Data", "Cidade", "Valor"]]
    linhas += [[dt.datetime(2026, 4, d), "Rio", float(d)] for d in range(1, 11)]
    linhas += [[], ["Resumo por cidade", None, None], ["Cidade", "Atendimentos", None], ["Rio", 10, None]]
    ds = servico.carregar(_xlsx({"P": linhas}), "duas.xlsx")
    assert len(ds.df) == 10
    assert "outro bloco" in _diag(ds)


def test_linha_em_branco_no_meio_nao_corta_a_tabela():
    linhas = [["Data", "Cidade", "Valor"]]
    linhas += [[dt.datetime(2026, 5, d), "Rio", float(d)] for d in range(1, 6)]
    linhas += [[]]
    linhas += [[dt.datetime(2026, 5, d), "Niterói", float(d)] for d in range(6, 11)]
    ds = servico.carregar(_xlsx({"P": linhas}), "lacuna.xlsx")
    assert len(ds.df) == 10


def test_total_sem_rotulo_na_ultima_linha():
    linhas = [["Cidade", "Qtd", "Valor"]] + [[c, q, v] for c, q, v in
                                              [("Rio", 2, 10.0), ("Niterói", 3, 20.0), ("Magé", 1, 5.0), ("Maricá", 4, 15.0)]]
    linhas.append([None, 10, 50.0])
    ds = servico.carregar(_xlsx({"P": linhas}), "total.xlsx")
    assert len(ds.df) == 4
    assert ds.df["Valor"].sum() == 50.0


def test_categoria_chamada_total_nao_e_descartada():
    """'Total'/'Parcial' como VALOR de uma coluna (tipo de glosa) é dado, não linha de total."""
    linhas = [["Tipo de glosa", "Cidade", "Valor"]]
    linhas += [["Total" if i % 2 else "Parcial", "Rio", float(i)] for i in range(1, 13)]
    ds = servico.carregar(_xlsx({"P": linhas}), "glosas.xlsx")
    assert len(ds.df) == 12
    assert (ds.df["Tipo de glosa"] == "Total").sum() == 6


def test_coluna_agrupada_e_preenchida_para_baixo():
    linhas = [["Região", "Cidade", "Valor"]]
    for regiao, cidades in (("Serrana", ["Petrópolis", "Teresópolis", "Friburgo"]),
                            ("Lagos", ["Cabo Frio", "Búzios"]), ("Metro", ["Rio", "Niterói", "Maricá"])):
        for i, c in enumerate(cidades):
            linhas.append([regiao if i == 0 else None, c, 10.0])
    ds = servico.carregar(_xlsx({"P": linhas}), "agrupado.xlsx")
    assert ds.df["Região"].isna().sum() == 0
    assert (ds.df["Região"] == "Serrana").sum() == 3


def test_datas_mistas_nativas_e_texto():
    linhas = [["Data", "Valor"]] + [[dt.datetime(2026, 1, d), 1.0] for d in range(1, 11)]
    linhas += [[f"{d:02d}/01/2026", 1.0] for d in range(11, 21)]
    ds = servico.carregar(_xlsx({"P": linhas}), "mistas.xlsx")
    assert _tipos(ds)["Data"] == DATA
    assert ds.df["Data"].notna().all()


def test_data_como_numero_serial_do_excel():
    linhas = [["Data de Emissão", "Valor"]] + [[45658 + d, 1.0] for d in range(10)]  # 45658 = 01/01/2025
    ds = servico.carregar(_xlsx({"P": linhas}), "serial.xlsx")
    assert _tipos(ds)["Data de Emissão"] == DATA
    assert ds.df["Data de Emissão"].iloc[0] == pd.Timestamp("2025-01-01")


def test_formato_largo_e_sinalizado():
    linhas = [["Produto", "Jan/26", "Fev/26", "Mar/26", "Abr/26"]]
    linhas += [[f"P{i}", i, i + 1, i + 2, i + 3] for i in range(1, 8)]
    ds = servico.carregar(_xlsx({"P": linhas}), "largo.xlsx")
    assert "Formato «largo»" in _diag(ds)


def test_planilha_limpa_nao_gera_diagnostico(csv_ptbr):
    ds = servico.carregar(csv_ptbr, "base.csv")
    assert ds.meta.diagnostico == []
    assert ds.meta.confianca >= 0.7
