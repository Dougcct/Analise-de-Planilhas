"""Leitor de relatórios em MATRIZ conta × mês (ex.: FATURAMENTO/DRE/CUSTOS do workbook HUNF).

Essas abas não são tabelas "uma linha por registro": são matrizes em que cada linha é uma conta
(ou convênio) e cada coluna é um mês. Este leitor "despivota" para a forma longa que o motor de
análise entende — uma linha por conta × mês —, dando de graça a evolução mensal, as quebras por
conta/grupo e as somas por recorte.

É dirigido pela seção `[leitura]` do modelo (`.toml`), então adaptar a outra planilha em matriz é
editar o modelo, sem tocar no código:

    [leitura]
    formato       = "matriz_mensal"
    coluna_rotulo = 3            # coluna (base 0) com o nome da conta/convênio
    rotulo_conta  = "Conta"      # nome da coluna de saída com o rótulo da linha
    rotulo_valor  = "Custo"      # nome da coluna de saída com o valor
    rotulo_grupo  = "Grupo"      # nome da coluna de grupo (se agrupamento != "nenhum")
    agrupamento   = "secoes"     # nenhum | codigo | maiusculas | secoes (ver abaixo)
    a_partir_de   = "2024-01"    # ignora os meses anteriores
    pular         = ["TOTAL"]    # rótulos (contém, maiúsculo) que são totais → não entram como dado
    linha_inicio  = 2            # 1ª e última linha da planilha (base 1) do bloco de dados
    linha_fim     = 133

    [[leitura.secao]]            # (agrupamento = "secoes") — cada bloco, pelo rótulo que o abre
    inicio = "CUSTOS HOSPITALARES"
    grupo  = "Custos hospitalares"
    [[leitura.secao]]            # bloco que só tem a linha de total: a própria linha é o dado
    inicio = "TOTAL DESPESAS SEGUROS DE BENS E IMOVEIS"
    grupo  = "Seguros de bens e imóveis"
    linha_unica = true

    [[leitura.referencia]]       # linha de total/contábil da planilha: guardada à parte (conferência)
    rotulo = "TOTAL COMPLEXO HOSPITALAR"
    nome   = "Total complexo hospitalar"

`agrupamento`:
- "nenhum"     — toda linha com rótulo é uma conta; não há coluna de grupo (ex.: FATURAMENTO);
- "codigo"     — conta termina em código numérico; as demais são cabeçalho/subtotal de grupo;
- "maiusculas" — linha em CAIXA ALTA é cabeçalho/subtotal de grupo; caixa mista é conta (ex.: DRE);
                 `contas = ["EPI"]` força rótulos em caixa alta a serem conta;
- "secoes"     — os blocos são declarados em `[[leitura.secao]]` (o mais robusto: segue a planilha).

Subtotais/cabeçalhos nunca entram como dados (não contam em dobro): viram a coluna de grupo e são
guardados em `meta.extras["subtotais"]` para a conferência «soma das contas × total da planilha».
Linhas de referência vão para `meta.extras["referencias"]`. Meses sem nenhum lançamento por conta
(futuros ou ainda não preenchidos) saem do detalhe e são informados no diagnóstico.

O leitor usa os VALORES que o Excel calculou e salvou (`data_only`) — as abas se referenciam por
fórmula; ver DOCS/13.
"""
from __future__ import annotations

import io
import re

import pandas as pd

from . import config, graficos
from .modelos import MetaLeitura
from .normalizacao import chave_texto, eh_data_nativa, texto_para_numero

MESES_MIN = 2                 # nº mínimo de datas numa linha para ela ser o cabeçalho de meses
_RE_CODIGO = re.compile(r"\s*\d{3,}\s*$")
_MESES_PT = ["jan", "fev", "mar", "abr", "mai", "jun", "jul", "ago", "set", "out", "nov", "dez"]


def config_modelo(modelo) -> dict | None:
    """Seção [leitura] do modelo se ele declara o formato matriz_mensal; senão None."""
    if modelo is None:
        return None
    cfg = modelo.regras.get("leitura", {})
    return cfg if cfg.get("formato") == "matriz_mensal" else None


def rotulo_mes(m: pd.Timestamp) -> str:
    return f"{_MESES_PT[m.month - 1]}/{str(m.year)[2:]}"


def _num(v):
    import numpy as np

    n = texto_para_numero(v)
    return None if n is None or (isinstance(n, float) and np.isnan(n)) else float(n)


def _mes(texto) -> pd.Timestamp:
    t = str(texto).strip()
    return pd.Timestamp(t + "-01" if len(t) == 7 else t).normalize().replace(day=1)


def _grade(dados: bytes, aba: str) -> list[tuple]:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(dados), data_only=True, read_only=True)
    try:
        if aba not in wb.sheetnames:
            raise ValueError(f"A aba «{aba}» não existe nesta planilha.")
        ws = wb[aba]
        ws.reset_dimensions()
        return [tuple(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()


def _linha_de_meses(linhas: list[tuple]) -> tuple[int, list[tuple[int, pd.Timestamp]]]:
    """Linha-cabeçalho dos meses (a que tem mais datas) e as colunas de mês (índice, 1º dia do mês)."""
    melhor, melhor_cols = -1, []
    for i, r in enumerate(linhas[:10]):
        cols = [(j, v) for j, v in enumerate(r) if eh_data_nativa(v)]
        if len(cols) > len(melhor_cols):
            melhor, melhor_cols = i, cols
    if len(melhor_cols) < MESES_MIN:
        raise ValueError("Não encontrei uma linha de cabeçalho com meses (datas) nesta aba.")
    meses = [(j, pd.Timestamp(pd.Timestamp(v).year, pd.Timestamp(v).month, 1)) for j, v in melhor_cols]
    return melhor, meses


def _eh_maiuscula(label: str) -> bool:
    return any(c.isalpha() for c in label) and label == label.upper()


def _grupo_ou_conta(label: str, modo: str) -> str:
    if modo == "codigo":
        return "conta" if _RE_CODIGO.search(label) else "grupo"
    if modo == "maiusculas":
        return "grupo" if _eh_maiuscula(label) else "conta"
    return "conta"


def ler(dados: bytes, nome_arquivo: str, aba: str, cfg: dict) -> tuple[pd.DataFrame, MetaLeitura]:
    """Converte a aba em matriz numa tabela longa (uma linha por conta × mês) + metadados."""
    linhas = _grade(dados, aba)
    i_cab, meses = _linha_de_meses(linhas)
    antes = 0
    if cfg.get("a_partir_de"):
        lim = _mes(cfg["a_partir_de"])
        antes = sum(1 for _, m in meses if m < lim)
        meses = [(j, m) for j, m in meses if m >= lim]
        if not meses:
            raise ValueError(f"A aba «{aba}» não tem meses a partir de {lim:%m/%Y}.")

    modo = str(cfg.get("agrupamento", "nenhum"))
    col_rot = int(cfg.get("coluna_rotulo", 0))
    k_conta = str(cfg.get("rotulo_conta", "Conta"))
    k_valor = str(cfg.get("rotulo_valor", "Valor"))
    k_mes = str(cfg.get("coluna_mes", "Mês"))
    k_grupo = str(cfg.get("rotulo_grupo", "Grupo"))
    limpar_codigo = bool(cfg.get("limpar_codigo", modo == "codigo"))
    pular = [str(p).upper() for p in cfg.get("pular", ["TOTAL"])]
    prefixos = [str(p) for p in cfg.get("pular_prefixos", ["*", "("])]
    # cabeçalhos de seção reais (prefixos, maiúsculo): com a lista, uma linha "de grupo" que NÃO casa
    # é tratada como conta — evita que uma conta sem código (ex.: «REPOUSO REMUNERADO») vire grupo
    grupos_pref = [str(p).upper() for p in cfg.get("grupos_prefixos", [])]
    forcadas = {chave_texto(x) for x in cfg.get("contas", [])}
    secoes = {chave_texto(s["inicio"]): s for s in cfg.get("secao", [])}
    refs_cfg = {chave_texto(r["rotulo"]): (i, r) for i, r in enumerate(cfg.get("referencia", []))}
    ini = int(cfg.get("linha_inicio", i_cab + 2))     # padrão: logo abaixo do cabeçalho
    fim = int(cfg.get("linha_fim", len(linhas)))
    usa_grupo = modo != "nenhum"

    registros: list[dict] = []
    unicas: dict[str, tuple[int, list[dict]]] = {}   # seção de linha única → (linha, registros); vale a última
    refs: dict[str, dict] = {}
    subtot: dict[str, dict] = {}                      # grupo → {ordem, linha, valores}
    grupo_atual, comecou = None, False
    n_grupos = n_pulos = n_datas = 0
    repetidas: list[str] = []

    def valores(r, contar_datas: bool) -> dict:
        nonlocal n_datas
        out = {}
        for j, mes in meses:
            v = r[j] if j < len(r) else None
            if eh_data_nativa(v):
                n_datas += contar_datas
                continue
            x = _num(v)
            if x is not None:
                out[mes] = x
        return out

    def subtotal(grupo, vals, nlin):
        """Candidato a total do grupo na planilha: fica o que tiver mais meses preenchidos."""
        if grupo is None or not vals:
            return
        atual = subtot.get(grupo)
        if atual is None:
            subtot[grupo] = {"ordem": len(subtot), "linha": nlin, "valores": vals}
        elif len(vals) > len(atual["valores"]):
            atual.update(linha=nlin, valores=vals)

    def emitir(destino: list, conta: str, vals: dict, nlin: int):
        for mes, val in vals.items():
            linha = {k_conta: conta, k_mes: mes, k_valor: val, "_lin": nlin}
            if usa_grupo:
                linha[k_grupo] = grupo_atual
            destino.append(linha)

    for p, r in enumerate(linhas):
        nlin = p + 1                                   # nº real da linha na planilha (base 1)
        if nlin < ini:
            continue
        if nlin > fim:
            break
        label = r[col_rot] if col_rot < len(r) else None
        if not isinstance(label, str) or not label.strip():
            if comecou and modo == "nenhum":
                break                                  # fim do bloco contíguo
            continue
        label = label.strip()
        k = chave_texto(label)

        if k in refs_cfg:                              # total/contábil da planilha: só conferência
            ordem, rc = refs_cfg[k]
            refs[rc.get("nome", label)] = {"ordem": ordem, "papel": rc.get("papel", ""), "linha": nlin,
                                           "valores": valores(r, False)}
            if modo == "maiusculas" and _eh_maiuscula(label) and k not in forcadas:
                grupo_atual = graficos.rotulo(label)
                subtotal(grupo_atual, refs[rc.get("nome", label)]["valores"], nlin)
            continue
        if any(pref and label.startswith(pref) for pref in prefixos):
            n_pulos += 1
            continue

        if modo == "secoes":
            if k in secoes:
                s = secoes[k]
                grupo_atual = s.get("grupo", graficos.rotulo(label))
                n_grupos += 1
                if s.get("linha_unica"):
                    regs: list[dict] = []
                    emitir(regs, s.get("conta", grupo_atual), valores(r, True), nlin)
                    if grupo_atual in unicas:
                        repetidas.append(f"«{label}» (linhas {unicas[grupo_atual][0]} e {nlin})")
                    unicas[grupo_atual] = (nlin, regs)
                else:
                    subtotal(grupo_atual, valores(r, False), nlin)
                continue
            if any(t in label.upper() for t in pular):
                subtotal(grupo_atual, valores(r, False), nlin)
                n_pulos += 1
                continue
            emitir(registros, label, valores(r, True), nlin)
            comecou = True
            continue

        if any(t in label.upper() for t in pular):     # linha de total/média
            if usa_grupo:
                subtotal(grupo_atual, valores(r, False), nlin)
            n_pulos += 1
            continue
        eh_grupo = k not in forcadas and _grupo_ou_conta(label, modo) == "grupo"
        if eh_grupo and grupos_pref and not any(label.upper().startswith(p_) for p_ in grupos_pref):
            eh_grupo = False                           # não é cabeçalho de seção: trata como conta
        if eh_grupo:
            grupo_atual = graficos.rotulo(label)       # "CUSTOS HOSPITALARES" → "Custos hospitalares"
            subtotal(grupo_atual, valores(r, False), nlin)
            n_grupos += 1
            continue
        conta = _RE_CODIGO.sub("", label).strip() if limpar_codigo else label
        emitir(registros, conta, valores(r, True), nlin)
        comecou = True

    for _, regs in unicas.values():
        registros += regs
    if not registros:
        raise ValueError(
            f"A aba «{aba}» está no formato matriz, mas nenhuma conta com valores foi lida no intervalo de "
            f"linhas {ini}–{fim}. Se as abas da planilha se referenciam por fórmula, abra o arquivo no Excel "
            "e salve (para gravar os valores calculados em cache); confira também linha_inicio/linha_fim e "
            "coluna_rotulo no modelo.")
    df = pd.DataFrame(registros)
    if limpar_codigo and modo == "secoes":
        df[k_conta] = df[k_conta].map(lambda s: _RE_CODIGO.sub("", s).strip() or s)

    # meses sem nenhum lançamento por conta (futuros/não preenchidos) saem do detalhe
    por_mes = df.groupby(k_mes)[k_valor].apply(lambda s: s.abs().sum())
    com_detalhe = {m for m, v in por_mes.items() if v}
    sem_detalhe = [m for _, m in meses if m not in com_detalhe]
    df = df[df[k_mes].isin(com_detalhe)]
    so_totais = [m for m in sem_detalhe
                 if any(x.get("valores", {}).get(m) for x in list(refs.values()) + list(subtot.values()))]
    validos = com_detalhe | set(so_totais)

    ordem = ([k_grupo] if usa_grupo else []) + [k_conta, k_mes, k_valor]
    df = df.set_index("_lin")[ordem]

    def tabela_extra(itens: dict, chave: str, com_papel: bool) -> pd.DataFrame:
        linhas_ = []
        for nome, info in itens.items():
            for mes, val in info["valores"].items():
                if mes in validos:
                    linha = {chave: nome, "Ordem": info["ordem"], k_mes: mes, "Valor": val,
                             "Linha na planilha": info["linha"]}
                    if com_papel:
                        linha["Papel"] = info.get("papel", "")
                    linhas_.append(linha)
        return pd.DataFrame(linhas_)

    contas = df[k_conta].nunique()
    meses_df = sorted(com_detalhe)
    meta = MetaLeitura(arquivo=nome_arquivo, aba=aba, abas=[aba], linha_cabecalho=i_cab,
                       linhas_cabecalho=[i_cab], confianca=1.0)
    meta.extras = {
        "referencias": tabela_extra(refs, "Linha", True),
        "subtotais": tabela_extra(subtot, k_grupo, False) if usa_grupo else pd.DataFrame(),
        "meses_so_totais": so_totais,
        "coluna_mes": k_mes, "coluna_valor": k_valor, "coluna_grupo": k_grupo if usa_grupo else None,
        "coluna_conta": k_conta,
    }
    meta.diagnostico = [
        f"Aba «{aba}» lida como matriz conta × mês: {contas} contas × {len(meses_df)} meses "
        f"({rotulo_mes(meses_df[0])} a {rotulo_mes(meses_df[-1])}) → {len(df)} linhas (uma por conta × mês).",
    ]
    if antes:
        meta.diagnostico.append(f"Base a partir de {rotulo_mes(meses[0][1])}: {antes} mês(es) anteriores "
                                "da planilha ficaram fora da análise.")
    futuros = [m for m in sem_detalhe if m not in so_totais]
    if futuros:
        meta.diagnostico.append(f"Meses sem nenhum lançamento ({rotulo_mes(futuros[0])} a {rotulo_mes(futuros[-1])}, "
                                f"{len(futuros)} mês(es)) foram ignorados.")
    if so_totais:
        meta.diagnostico.append("Mês(es) com só os totais da planilha, sem o detalhe por conta: "
                                + ", ".join(rotulo_mes(m) for m in so_totais)
                                + " — ficam fora da análise detalhada e aparecem na conferência.")
    if n_grupos:
        meta.diagnostico.append(f"{n_grupos} linha(s) de grupo/subtotal viraram a coluna «{k_grupo}» "
                                "(não entram como dados, para não contar em dobro).")
    if n_pulos:
        meta.diagnostico.append(f"{n_pulos} linha(s) de total/anotação foram ignoradas como dado.")
    if refs:
        meta.diagnostico.append("Linhas de total da planilha guardadas para conferência (não somam no detalhe): "
                                + ", ".join(f"«{n}»" for n in refs) + ".")
    if repetidas:
        meta.diagnostico.append("Linha repetida — valeu a última: " + "; ".join(repetidas) + ".")
    if n_datas:
        meta.diagnostico.append(f"{config.PREFIXO_AVISO}{n_datas} célula(s) com DATA digitada no lugar de valor foram "
                                "ignoradas — conferir na planilha.")
    return df, meta
